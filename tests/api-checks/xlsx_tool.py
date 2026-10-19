"""Dung va doc file .xlsx cho bo kiem excel-import.sh (can openpyxl).

Dung de sinh file nhap lieu thu, kem vai file "xau" (bom zip, file ngau nhien) ma API phai tu choi.

  python xlsx_tool.py make OUT --header "A|B|C" [--sheet TEN] [--pad N] [--decoy] < dong
      Moi dong stdin (UTF-8) la mot dong Excel, cac o cach nhau bang dau gach dung. Quy uoc trong o:
        #123     ghi la SO (de thu o so mat so 0 dau, vd so dien thoai)
        ~text    chuan hoa NFD (thu chuoi tieng Viet dang to hop)
        \\uXXXX  ma Unicode, de nguon bo kiem co the toan ASCII
      --pad N    them N dong tieu de truoc hang tieu de cot (file thuc te thuong co dong ten truong)
      --decoy    dat mot trang "Huong dan" len truoc trang du lieu

  python xlsx_tool.py fill TEMPLATE OUT < dong
      Mo file mau do API tai ve, ghi cac dong vao trang dau tu dong 2 — dung lai dung luong nguoi dung lam that.

  python xlsx_tool.py info FILE
      In JSON: ten cac trang, hang tieu de cua trang dau, so o co danh sach chon, noi dung trang "Danh sach".

  python xlsx_tool.py dump FILE
      In JSON noi dung trang dau: ten cac trang, cot/hang dang co dinh, tung dong (o trong la null), so o la CONG THUC,
      so o la SO. Dung doc lai bang diem server xuat ra.

  python xlsx_tool.py rows OUT --header "A|B" --count N --template "ZZ{i}|Ten {i}"
      Sinh N dong tu mau, {i} la so thu tu tu 1.

  python xlsx_tool.py bomb SRC OUT --mb 120
      Chep SRC sang OUT va them mot muc zip nen duoc (toan so 0) phinh ra MB megabyte khi giai nen.

  python xlsx_tool.py junk OUT --mb 6
      Ghi mot file ngau nhien OUT dung MB megabyte (thu tran dung luong).
"""
import argparse
import json
import os
import sys
import unicodedata
import zipfile

from openpyxl import Workbook, load_workbook


def decode_cell(raw: str):
    if raw.startswith("#"):
        text = raw[1:]
        return int(text) if text.lstrip("-").isdigit() else float(text)
    nfd = raw.startswith("~")
    if nfd:
        raw = raw[1:]
    value = raw.encode("ascii").decode("unicode_escape") if "\\u" in raw else raw
    return unicodedata.normalize("NFD", value) if nfd else value


def read_rows() -> list[list]:
    text = sys.stdin.buffer.read().decode("utf-8")
    rows = []
    for line in text.splitlines():
        if line.strip() == "":
            continue
        rows.append([decode_cell(cell) for cell in line.split("|")])
    return rows


def write_book(out: str, header: list[str], rows: list[list], sheet: str, pad: int, decoy: bool) -> None:
    book = Workbook()
    active = book.active
    if decoy:
        active.title = "Huong dan"
        active.append(["Day la trang huong dan, khong phai du lieu"])
        active = book.create_sheet(sheet)
    else:
        active.title = sheet

    for index in range(pad):
        active.append([f"Danh sach nhap lieu (dong tieu de {index + 1})"])
    active.append(header)
    for row in rows:
        active.append(row)
    book.save(out)


def command_make(args) -> None:
    header = [decode_cell(cell) for cell in args.header.split("|")]
    write_book(args.out, header, read_rows(), args.sheet, args.pad, args.decoy)


def command_fill(args) -> None:
    book = load_workbook(args.template)
    sheet = book.worksheets[0]
    for offset, row in enumerate(read_rows()):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=2 + offset, column=column, value=value)
    book.save(args.out)


def command_info(args) -> None:
    book = load_workbook(args.file)
    first = book.worksheets[0]
    validations = sum(len(list(v.sqref.ranges)) for v in first.data_validations.dataValidation)
    lists = {}
    if "Danh sách" in book.sheetnames:
        sheet = book["Danh sách"]
        for column in sheet.iter_cols(min_row=1):
            values = [cell.value for cell in column if cell.value is not None]
            if values:
                lists[values[0]] = values[1:]
    info = {
        "sheets": book.sheetnames,
        "header": [cell.value for cell in first[1] if cell.value is not None],
        "validations": validations,
        "lists": lists,
        "guideLines": [row[0].value for row in book["Hướng dẫn"].iter_rows()] if "Hướng dẫn" in book.sheetnames else [],
    }
    sys.stdout.buffer.write(json.dumps(info, ensure_ascii=False).encode("utf-8"))


def command_dump(args) -> None:
    book = load_workbook(args.file)
    sheet = book.worksheets[0]
    rows = []
    formulas = 0
    numbers = 0
    for row in sheet.iter_rows():
        values = []
        for cell in row:
            if cell.data_type == "f":
                formulas += 1
            elif cell.data_type == "n" and cell.value is not None:
                numbers += 1
            values.append(cell.value)
        rows.append(values)
    dump = {
        "sheets": book.sheetnames,
        "freeze": sheet.freeze_panes,
        "rows": rows,
        "formulas": formulas,
        "numbers": numbers,
    }
    sys.stdout.buffer.write(json.dumps(dump, ensure_ascii=False).encode("utf-8"))


def command_rows(args) -> None:
    header = [decode_cell(cell) for cell in args.header.split("|")]
    rows = [
        [decode_cell(cell.replace("{i}", str(index))) for cell in args.template.split("|")]
        for index in range(1, args.count + 1)
    ]
    write_book(args.out, header, rows, args.sheet, 0, False)


def command_bomb(args) -> None:
    with zipfile.ZipFile(args.src) as source, zipfile.ZipFile(args.out, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            target.writestr(item, source.read(item.filename))
        chunk = b"\0" * (1024 * 1024)
        with target.open("xl/junk.bin", "w") as junk:
            for _ in range(args.mb):
                junk.write(chunk)


def command_junk(args) -> None:
    with open(args.out, "wb") as handle:
        for _ in range(args.mb):
            handle.write(os.urandom(1024 * 1024))


def main() -> None:
    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="command", required=True)

    make = sub.add_parser("make")
    make.add_argument("out")
    make.add_argument("--header", required=True)
    make.add_argument("--sheet", default="Sheet1")
    make.add_argument("--pad", type=int, default=0)
    make.add_argument("--decoy", action="store_true")
    make.set_defaults(run=command_make)

    fill = sub.add_parser("fill")
    fill.add_argument("template")
    fill.add_argument("out")
    fill.set_defaults(run=command_fill)

    info = sub.add_parser("info")
    info.add_argument("file")
    info.set_defaults(run=command_info)

    dump = sub.add_parser("dump")
    dump.add_argument("file")
    dump.set_defaults(run=command_dump)

    rows = sub.add_parser("rows")
    rows.add_argument("out")
    rows.add_argument("--header", required=True)
    rows.add_argument("--count", type=int, required=True)
    rows.add_argument("--template", required=True)
    rows.add_argument("--sheet", default="Sheet1")
    rows.set_defaults(run=command_rows)

    bomb = sub.add_parser("bomb")
    bomb.add_argument("src")
    bomb.add_argument("out")
    bomb.add_argument("--mb", type=int, default=120)
    bomb.set_defaults(run=command_bomb)

    junk = sub.add_parser("junk")
    junk.add_argument("out")
    junk.add_argument("--mb", type=int, default=6)
    junk.set_defaults(run=command_junk)

    args = parser.parse_args()
    args.run(args)


if __name__ == "__main__":
    main()
